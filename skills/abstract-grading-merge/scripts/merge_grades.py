#!/usr/bin/env python3
"""
Merge multiple reviewers' abstract-grading Excel files (QM2027-style: one
file per reviewer, each with an 'Oral' sheet and usually a 'Poster' sheet,
column layouts that differ reviewer to reviewer) into one workbook with a
per-abstract average score, a unioned "Merge with" suggestion, and
per-reviewer comments stacked as "Name: ...".

Usage:
  python3 merge_grades.py <folder> [--track NAME] [--out PATH]

<folder>   directory containing the reviewers' .xlsx files (and nothing
           else that should be ingested -- lock files '~$*.xlsx' and any
           previous '*_merged_*.xlsx' output are skipped automatically).
--track    a token (e.g. "Small_Systems") stripped from filenames when
           guessing each reviewer's display name. If omitted, inferred
           from the most common prefix across filenames.
--out      output path. Default: "<folder>/<track>_merged_<YYYYMMDD>.xlsx"
           (never overwrites an existing snapshot -- a numeric suffix is
           added if the name collides).

See SKILL.md for the full write-up of the column-detection heuristics,
the evenness check, and what to do when a new reviewer's file doesn't
parse cleanly.
"""
import argparse
import datetime
import glob
import os
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = "DDEBF7"
CONFLICT_COLOR = "CC0000"

MERGE_PAT = re.compile(
    r"merge\s*(?:with|w/|w\.)?\s*#?(\d+)(?:\s*(?:or|,|and)\s*#?(\d+))?", re.I
)

ID_ALIASES = {"id", "abstract #", "abstract id", "#id", "abstract"}
TITLE_ALIASES = {"title"}
AFFIL_ALIASES = {"affiliation"}
CATEGORY_ALIASES = {"category", "type"}
COLLAB_ALIASES = {
    "collaboration",
    "if yes, please enter the collaboration name.",
}
SCORE_ALIASES = {"score", "grade"}
COMMENT_ALIASES = {"comments", "comment", "note", "notes"}
MERGEWITH_ALIASES = {"merge with"}
SKIP_SHEETS = {"summary", "evenness check", "evenness"}


def norm(h):
    return (h or "").strip().lower()


def find_col(headers, aliases):
    for i, h in enumerate(headers):
        if norm(h) in aliases:
            return i
    return None


def read_sheet_rows(ws):
    """Read header row + data rows, stopping at the first fully-blank row
    (reviewers often leave a breakdown/summary table below the real data,
    e.g. a per-category or per-collaboration grade-count block)."""
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in vals):
            break
        rows.append(vals)
    return headers, rows


def parse_merge_targets(text):
    targets = set()
    if not text:
        return targets
    for m in MERGE_PAT.finditer(text):
        for g in m.groups():
            if g:
                targets.add(int(g))
    return targets


def guess_label(path, track_token):
    stem = os.path.splitext(os.path.basename(path))[0]
    s = stem
    if track_token:
        s = re.sub(re.escape(track_token), "", s, flags=re.I)
    s = re.sub(r"^[\d_]+", "", s)  # leading track-number prefix like "17_"
    s = re.sub(r"(grading|scores|score|grades|grade)", "", s, flags=re.I)
    s = s.strip("_- ")
    return s or stem


def detect_track_token(xlsx_paths):
    stems = [os.path.splitext(os.path.basename(p))[0] for p in xlsx_paths]
    tokens = [re.sub(r"^\d+_", "", s) for s in stems]
    counts = Counter()
    for t in tokens:
        parts = t.split("_")
        for n in (2, 3):
            if len(parts) >= n:
                counts["_".join(parts[:n])] += 1
    if counts:
        return counts.most_common(1)[0][0]
    return ""


def list_input_files(folder):
    paths = sorted(glob.glob(os.path.join(folder, "*.xlsx")))
    out = []
    for p in paths:
        base = os.path.basename(p)
        if base.startswith("~$"):
            continue
        if re.search(r"_merged_\d{8}", base):
            continue
        out.append(p)
    return out


def extract_sheet(path, ws, label):
    """Returns list of dicts: id, title, affiliation, category, collaboration,
    score, comment, merge_targets."""
    headers, rows = read_sheet_rows(ws)
    idc = find_col(headers, ID_ALIASES)
    if idc is None:
        return []
    tc = find_col(headers, TITLE_ALIASES)
    afc = find_col(headers, AFFIL_ALIASES)
    catc = find_col(headers, CATEGORY_ALIASES)
    colc = find_col(headers, COLLAB_ALIASES)
    scc = find_col(headers, SCORE_ALIASES)
    cmc = find_col(headers, COMMENT_ALIASES)
    mgc = find_col(headers, MERGEWITH_ALIASES)

    # Indico-raw-export quirk (seen in QM2027 "YL"-style files): an
    # unlabeled free-note column can sit immediately after "Score", right
    # before "Score standard deviation". Treat it as an extra comment
    # source when present and no named comment column was found.
    if cmc is None and scc is not None:
        sdc = find_col(headers, {"score standard deviation"})
        if sdc is not None and sdc - scc == 2:
            cmc = scc + 1

    out = []
    for row in rows:
        idv = row[idc] if idc < len(row) else None
        if not isinstance(idv, int):
            continue
        rec = {
            "id": idv,
            "title": row[tc] if tc is not None else None,
            "affiliation": row[afc] if afc is not None else None,
            "category": row[catc] if catc is not None else None,
            "collaboration": row[colc] if colc is not None else None,
            "score": row[scc] if scc is not None else None,
            "comment": (str(row[cmc]).strip() if cmc is not None and row[cmc] else None),
            "merge_targets": set(),
        }
        if mgc is not None and row[mgc] is not None:
            try:
                rec["merge_targets"].add(int(row[mgc]))
            except (ValueError, TypeError):
                pass
        if rec["comment"]:
            rec["merge_targets"] |= parse_merge_targets(rec["comment"])
        out.append(rec)
    return out


def evenness(score_counter, n):
    target = n / 4 if n else 0
    dist = [score_counter.get(k, 0) for k in (3, 2, 1, 0)]
    deviation = max(abs(d - target) for d in dist) if n else 0
    even = deviation <= max(2, target * 0.25)
    return dist, target, even


def build_sheet_records(files_by_label, sheet_name):
    """files_by_label: {label: path}. Returns (records dict keyed by id,
    graders actually present for this sheet, evenness report list)."""
    records = {}
    graders_present = []
    report = []
    for label, path in files_by_label.items():
        wb = openpyxl.load_workbook(path, data_only=True)
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = extract_sheet(path, ws, label)
        if not rows:
            continue
        c = Counter()
        n = 0
        for rec in rows:
            idv = rec["id"]
            master = records.setdefault(
                idv,
                {
                    "id": idv,
                    "title": None,
                    "affiliation": None,
                    "category": None,
                    "collaboration": None,
                    "scores": {},
                    "comments": {},
                    "merge_targets": set(),
                },
            )
            for k in ("title", "affiliation", "category", "collaboration"):
                if not master[k] and rec[k]:
                    master[k] = rec[k]
            if rec["score"] is not None:
                master["scores"][label] = rec["score"]
                if rec["score"] in (0, 1, 2, 3):
                    c[rec["score"]] += 1
                    n += 1
            if rec["comment"]:
                master["comments"][label] = rec["comment"]
            master["merge_targets"] |= rec["merge_targets"]
        if n > 0:
            graders_present.append(label)
        dist, target, even = evenness(c, n)
        report.append(
            {
                "label": label,
                "sheet": sheet_name,
                "n": n,
                "target": target,
                "dist": dist,
                "even": even,
            }
        )
    # symmetrize merge targets
    changed = True
    while changed:
        changed = False
        for idv, rec in records.items():
            for t in list(rec["merge_targets"]):
                if t in records and idv not in records[t]["merge_targets"]:
                    records[t]["merge_targets"].add(idv)
                    changed = True
    return records, graders_present, report


def poster_quality_ok(report, total_abstracts, min_graders=2, min_frac=0.8):
    """Heuristic: is there enough real per-abstract poster grading to
    build a full merged average table, or should we fall back to a
    0-score-only flag list? A grader counts as 'real' only if they
    scored most of the pool (min_frac of total_abstracts) AND didn't
    give every abstract the exact same score (a classic placeholder
    sign, e.g. a reviewer who hasn't actually graded yet)."""
    if total_abstracts == 0:
        return False
    real = 0
    for r in report:
        if r["n"] == 0:
            continue
        if r["n"] / total_abstracts < min_frac:
            continue  # too few of the pool scored -> effectively ungraded
        if max(r["dist"]) == r["n"] and r["n"] > 3:
            continue  # every scored row got the same value -> placeholder
        real += 1
    return real >= min_graders


def style_header(ws, headers):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)


def write_merged_sheet(wb, sheet_title, records, graders):
    ws = wb.create_sheet(sheet_title)
    headers = ["Id", "Title", "Affiliation", "Category", "Collaboration"] + graders + [
        "Average",
        "Merge with",
        "Comments",
    ]
    style_header(ws, headers)
    score_start = 6
    score_end = score_start + len(graders) - 1

    def avg_of(rec):
        vals = [rec["scores"][g] for g in graders if g in rec["scores"]]
        return sum(vals) / len(vals) if vals else -1

    sorted_ids = sorted(records.keys(), key=lambda k: (-avg_of(records[k]), k))
    row_of = {}
    r = 2
    for idv in sorted_ids:
        rec = records[idv]
        ws.cell(row=r, column=1, value=idv).font = Font(name=FONT)
        ws.cell(row=r, column=2, value=rec["title"]).font = Font(name=FONT)
        ws.cell(row=r, column=3, value=rec["affiliation"]).font = Font(name=FONT)
        ws.cell(row=r, column=4, value=rec["category"]).font = Font(name=FONT)
        ws.cell(row=r, column=5, value=rec["collaboration"]).font = Font(name=FONT)
        for gi, g in enumerate(graders):
            ws.cell(row=r, column=score_start + gi, value=rec["scores"].get(g)).font = Font(
                name=FONT
            )
        row_of[idv] = r
        r += 1

    for idv in sorted_ids:
        rec = records[idv]
        rr = row_of[idv]
        rng = f"{get_column_letter(score_start)}{rr}:{get_column_letter(score_end)}{rr}"
        ws.cell(row=rr, column=score_end + 1, value=f"=AVERAGE({rng})").font = Font(name=FONT)

        targets = sorted(rec["merge_targets"])
        merge_cell = ws.cell(row=rr, column=score_end + 2)
        comment_lines = []
        if targets:
            merge_cell.value = ", ".join(str(t) for t in targets)
            merge_cell.font = Font(
                name=FONT, color=CONFLICT_COLOR if len(targets) > 1 else "000000", bold=len(targets) > 1
            )
            if len(targets) > 1:
                comment_lines.append("⚠ Merge-with suggestions disagree across reviewers")

        for g in graders:
            c = rec["comments"].get(g)
            if c:
                comment_lines.append(f"{g}: {c}")
        cmt_cell = ws.cell(row=rr, column=score_end + 3)
        if comment_lines:
            cmt_cell.value = "\n".join(comment_lines)
            cmt_cell.alignment = Alignment(wrap_text=True, vertical="top")
        cmt_cell.font = Font(name=FONT)

    widths = {1: 6, 2: 45, 3: 28, 4: 10, 5: 18, score_end + 1: 9, score_end + 2: 12, score_end + 3: 55}
    for gi in range(len(graders)):
        widths[score_start + gi] = 9
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_zero_flag_sheet(wb, sheet_title, records, note):
    ws = wb.create_sheet(sheet_title)
    headers = ["Id", "Title", "Collaboration", "Reviewer", "Score", "Comment"]
    style_header(ws, headers)
    r = 2
    for idv, rec in sorted(records.items()):
        for g, sc in rec["scores"].items():
            if sc == 0:
                ws.cell(row=r, column=1, value=idv).font = Font(name=FONT)
                ws.cell(row=r, column=2, value=rec["title"]).font = Font(name=FONT)
                ws.cell(row=r, column=3, value=rec["collaboration"]).font = Font(name=FONT)
                ws.cell(row=r, column=4, value=g).font = Font(name=FONT)
                ws.cell(row=r, column=5, value=0).font = Font(name=FONT)
                ws.cell(row=r, column=6, value=rec["comments"].get(g, "")).font = Font(name=FONT)
                r += 1
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 18
    ws.cell(row=r + 1, column=1, value=note).font = Font(name=FONT, italic=True, size=9)
    return ws


def write_evenness_sheet(wb, reports):
    ws = wb.create_sheet("Evenness check")
    headers = ["Reviewer", "Sheet", "Graded", "Target/grade", "Grade 3", "Grade 2", "Grade 1", "Grade 0", "Even?"]
    style_header(ws, headers)
    r = 2
    for rep in reports:
        ws.cell(row=r, column=1, value=rep["label"]).font = Font(name=FONT)
        ws.cell(row=r, column=2, value=rep["sheet"]).font = Font(name=FONT)
        ws.cell(row=r, column=3, value=rep["n"]).font = Font(name=FONT)
        ws.cell(row=r, column=4, value=round(rep["target"], 2)).font = Font(name=FONT)
        for i, v in enumerate(rep["dist"]):
            ws.cell(row=r, column=5 + i, value=v).font = Font(name=FONT)
        cell = ws.cell(row=r, column=9, value="OK" if rep["even"] else "UNEVEN")
        cell.font = Font(name=FONT, bold=not rep["even"], color="CC0000" if not rep["even"] else "000000")
        r += 1
    for col, w in {1: 14, 2: 10}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("--track", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    files = list_input_files(args.folder)
    if not files:
        print("No input .xlsx files found.", file=sys.stderr)
        sys.exit(1)

    track_token = args.track or detect_track_token(files)
    labels = {}
    for p in files:
        lbl = guess_label(p, track_token)
        labels[lbl] = p
    print("Reviewer files detected:")
    for lbl, p in labels.items():
        print(f"  {lbl:12s} <- {os.path.basename(p)}")

    oral_records, oral_graders, oral_report = build_sheet_records(labels, "Oral")
    poster_records, poster_graders, poster_report = build_sheet_records(labels, "Poster")

    print("\nEvenness check:")
    for rep in oral_report + poster_report:
        flag = "OK" if rep["even"] else "UNEVEN"
        print(
            f"  {rep['label']:10s} {rep['sheet']:7s} n={rep['n']:3d} "
            f"target={rep['target']:5.2f} dist(3,2,1,0)={rep['dist']} {flag}"
        )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    if oral_records:
        write_merged_sheet(wb, "Oral", oral_records, oral_graders)

    poster_ok = poster_quality_ok(poster_report, len(poster_records))
    if poster_records and poster_ok:
        write_merged_sheet(wb, "Poster", poster_records, poster_graders)
    elif poster_records:
        note = (
            "Poster scores were not merged/averaged: fewer than 2 reviewers had "
            "complete, non-placeholder grading (see 'Evenness check' sheet). "
            "Only abstracts that any reviewer scored 0 are listed below; re-run "
            "once more reviewers' Poster scores are complete."
        )
        write_zero_flag_sheet(wb, "Poster_zero_flags", poster_records, note)
        print("\nPoster: insufficient real grading data -> wrote zero-flag list only.")

    write_evenness_sheet(wb, oral_report + poster_report)

    out = args.out
    if not out:
        today = datetime.datetime.now().strftime("%Y%m%d")
        base = f"{track_token or 'grading'}_merged_{today}.xlsx"
        out = os.path.join(args.folder, base)
        n = 1
        while os.path.exists(out):
            n += 1
            out = os.path.join(args.folder, f"{track_token or 'grading'}_merged_{today}_{n}.xlsx")

    wb.save(out)
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
